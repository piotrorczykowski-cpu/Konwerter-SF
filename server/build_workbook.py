import json
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY='153D6B'; BLUE='1F4E78'; PALE='DCE6F1'; LIGHT='F5F8FC'; GRID='B9C7D6'; MUTED='5C6773'
thin=Side(style='thin', color=GRID)

def value(row, key):
    v=row.get(key, '')
    return '' if v is None else v

def set_widths(ws, widths):
    for col, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(col)].width=width

def title(ws, heading, subtitle):
    ws.sheet_view.showGridLines=False
    ws['A1']=heading; ws['A1'].font=Font(name='Arial', bold=True, size=15, color=NAVY)
    ws['A2']=subtitle; ws['A2'].font=Font(name='Arial', italic=True, size=9, color=MUTED)
    ws.row_dimensions[1].height=27; ws.row_dimensions[2].height=19

def table(ws, start, headers, rows, keys, widths):
    for c, h in enumerate(headers,1):
        cell=ws.cell(start,c,h); cell.fill=PatternFill('solid',fgColor=NAVY); cell.font=Font(name='Arial',bold=True,size=9,color='FFFFFF'); cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.row_dimensions[start].height=30
    for idx, row in enumerate(rows,start+1):
        for c, key in enumerate(keys,1):
            v=value(row,key)
            if key=='Pozycja' and row.get('Poziom',0): v='   '*row['Poziom']+str(v)
            if isinstance(v,(int,float)) and ('udział' in key.lower() or row.get('Format')=='percent'): v=v/100
            section=row.get('Poziom')==0 and row.get('Pozycja')!='Suma'
            total=row.get('Pozycja')=='Suma'
            cell=ws.cell(idx,c,v)
            cell.font=Font(name='Arial',size=9,bold=(section or total),color=(NAVY if section else '000000'))
            cell.alignment=Alignment(horizontal=('left' if c==1 else 'right'),vertical='center',wrap_text=(c==1))
            cell.border=Border(bottom=thin)
            if section: cell.fill=PatternFill('solid',fgColor=PALE)
            elif idx%2==0: cell.fill=PatternFill('solid',fgColor=LIGHT)
            if total: cell.border=Border(top=Side(style='medium',color=BLUE),bottom=Side(style='double',color=BLUE))
            low=headers[c-1].lower()
            if isinstance(v,(int,float)):
                if 'udział' in low or 'procent' in low: cell.number_format='0.00%;(0.00%);-'
                elif 'wan' in low or ('wartość' in low and 'certyfikat' in low): cell.number_format='#,##0.00;(#,##0.00);-'
                else: cell.number_format='#,##0;(#,##0);-'
    end=start+len(rows)
    if rows:
        ref=f'A{start}:{get_column_letter(len(headers))}{end}'
        tab=Table(displayName=f'T{ws.title.replace(" ","")[:12]}{start}',ref=ref)
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False,showColumnStripes=False)
        ws.add_table(tab)
    set_widths(ws,widths)
    ws.freeze_panes='B5'
    return end

def financial(wb,name,heading,period,rows,extra,note,cy,py):
    ws=wb.create_sheet(name); title(ws,heading,period)
    end=table(ws,4,['Pozycja',cy,py],rows,['Pozycja',cy,py],[76,18,18])
    if extra:
        ws.cell(end+2,1,'Dane uzupełniające').font=Font(name='Arial',bold=True,size=10,color=NAVY)
        end=table(ws,end+3,['Pozycja',cy,py],extra,['Pozycja',cy,py],[76,18,18])
    ws.cell(end+2,1,note).font=Font(name='Arial',italic=True,size=8,color=MUTED)

def build(model_path, output_path):
    data=json.loads(Path(model_path).read_text(encoding='utf-8')); meta=data.get('meta',{})
    cy=str(meta.get('current_year','2024')); py=str(meta.get('comparative_year','2023')); start=meta.get('period_from',f'{cy}-01-01'); end=meta.get('period_to',f'{cy}-12-31')
    wb=Workbook(); wb.remove(wb.active)
    ws=wb.create_sheet('Wprowadzenie'); title(ws,str(meta.get('fund_name','Fundusz inwestycyjny')).upper(),f'Sprawozdanie finansowe za okres od {start} do {end}.')
    table(ws,4,['Pole','Wartość'],data['intro'],['Pole','Wartość'],[34,118])
    ws=wb.create_sheet('Zestawienie lokat'); title(ws,'ZESTAWIENIE LOKAT',f'Według stanu na {end}. Wszystkie dane w tys. zł, z wyjątkiem liczby instrumentów.')
    r=table(ws,4,['Składniki lokat',f'Cena nabycia {cy}',f'Wycena {cy}',f'Udział {cy}',f'Cena nabycia {py}',f'Wycena {py}',f'Udział {py}'],data['main_rows'],['Pozycja',f'{cy} cena nabycia',f'{cy} wycena',f'{cy} udział',f'{py} cena nabycia',f'{py} wycena',f'{py} udział'],[64,18,18,14,18,18,14])
    for heading,headers,keys,rows,widths in [
        ('TABELA UZUPEŁNIAJĄCA - AKCJE',['Nazwa lokaty','Rodzaj rynku','Nazwa rynku','Liczba','Kraj emitenta','Cena nabycia','Wycena','Udział'],['Nazwa lokaty','Rodzaj rynku','Nazwa rynku','Liczba','Kraj emitenta','Cena nabycia','Wycena','Udział'],data['holdings'],[55,22,20,14,18,16,16,14]),
        ('TABELA UZUPEŁNIAJĄCA - UDZIAŁY W SPÓŁKACH Z O.O.',['Nazwa udziału','Nazwa spółki','Liczba','Cena nabycia','Wycena','Udział'],['Nazwa udziału','Nazwa spółki','Liczba','Cena nabycia','Wycena','Udział'],data['equity'],[42,42,16,18,18,14]),
        ('TABELA UZUPEŁNIAJĄCA - NIERUCHOMOŚCI',['Data nabycia','Numer księgi wieczystej','Kraj położenia','Adres','Powierzchnia','Służebności','Cena nabycia','Wycena','Udział'],['Data nabycia','Numer księgi wieczystej','Kraj położenia','Adres','Powierzchnia','Służebności','Cena nabycia','Wycena','Udział'],data['properties'],[16,25,16,48,16,65,16,16,14])]:
        ws.cell(r+2,1,heading).font=Font(name='Arial',bold=True,size=10,color=NAVY); r=table(ws,r+3,headers,rows,keys,widths)
    financial(wb,'Bilans','BILANS',f'Stan na {end} i dane porównawcze za {py} r.',data['balance'],data['balance_extra'],'Wszystkie dane w tys. złotych za wyjątkiem liczby certyfikatów inwestycyjnych oraz WAN/CI.',cy,py)
    financial(wb,'Wynik z operacji','RACHUNEK WYNIKU Z OPERACJI',f'Za okres od {start} do {end} oraz dane porównawcze za {py} r.',data['income'],data['income_extra'],'Wszystkie dane w tys. złotych.',cy,py)
    financial(wb,'Zmiany aktywów netto','ZESTAWIENIE ZMIAN W AKTYWACH NETTO',f'Za okres od {start} do {end} oraz dane porównawcze za {py} r.',data['changes'],data.get('nav_changes'),'Wszystkie dane w tys. złotych.',cy,py)
    financial(wb,'Przepływy','RACHUNEK PRZEPŁYWÓW PIENIĘŻNYCH',f'Za okres od {start} do {end} oraz dane porównawcze za {py} r.',data['cashflows'],None,'Wszystkie dane w tys. złotych.',cy,py)
    for item in data.get('formatted_embedded',[]):
        ws=wb.create_sheet(item['Arkusz'][:31]); title(ws,item['Arkusz'].upper(),f'Pełna treść pliku {item["Plik"]}, podzielona na strony i sekcje.')
        table(ws,4,['Strona','Sekcja','Treść'],item['rows'],['Strona','Sekcja','Treść'],[12,45,95])
        for i,row in enumerate(item['rows'],5):
            text_len=len(str(row.get('Treść','')))
            ws.cell(i,3).alignment=Alignment(vertical='top',wrap_text=True); ws.row_dimensions[i].height=28 if row.get('Typ')=='Nagłówek' else min(110,max(36,18+text_len//95*14))
            if row.get('Typ')=='Nagłówek':
                for c in range(1,4): ws.cell(i,c).fill=PatternFill('solid',fgColor=PALE); ws.cell(i,c).font=Font(name='Arial',bold=True,size=9,color=NAVY)
    wb.save(output_path)

if __name__=='__main__': build(sys.argv[1],sys.argv[2])
